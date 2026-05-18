from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_DIR = Path(r"D:\New folder\Pos_System\manual_output")
OUTPUT_DIR.mkdir(exist_ok=True)


def set_cell_text(cell, text, rtl=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if rtl:
        set_paragraph_rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    run.font.name = "Tahoma" if rtl else "Segoe UI"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), run.font.name)
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_paragraph_rtl(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_run_font(run, font_name, size=None, bold=False, color=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_paragraph(doc, text="", style=None, rtl=False, bold=False):
    p = doc.add_paragraph(style=style)
    if rtl:
        set_paragraph_rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    set_run_font(run, "Tahoma" if rtl else "Segoe UI", 10.5, bold)
    return p


def add_heading(doc, text, level=1, rtl=False):
    p = doc.add_heading(level=level)
    if rtl:
        set_paragraph_rtl(p)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    size = {0: 22, 1: 17, 2: 14, 3: 12}.get(level, 12)
    set_run_font(run, "Tahoma" if rtl else "Segoe UI", size, True, (15, 23, 42))
    return p


def add_bullets(doc, items, rtl=False):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        if rtl:
            set_paragraph_rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(item)
        set_run_font(run, "Tahoma" if rtl else "Segoe UI", 10.5)


def add_numbered(doc, items, rtl=False):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        if rtl:
            set_paragraph_rtl(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(item)
        set_run_font(run, "Tahoma" if rtl else "Segoe UI", 10.5)


def add_table(doc, headers, rows, rtl=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT if rtl else WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, rtl)
        for run in table.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, rtl)
    doc.add_paragraph()


def setup_document(rtl=False):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    for style_name in ["Normal", "List Bullet", "List Number"]:
        style = doc.styles[style_name]
        style.font.name = "Tahoma" if rtl else "Segoe UI"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma" if rtl else "Segoe UI")
        style.font.size = Pt(10.5)
    return doc


english_sections = [
    ("1. Purpose Of The Application", [
        "This POS System manages sales, stock, customers, suppliers, balances, reports, users, permissions, settings, and work history from one dashboard.",
        "The system is designed for daily cashier work, inventory control, supplier/customer balance tracking, and management reporting.",
        "Access to each screen and action depends on the permissions granted in Settings > Options."
    ]),
    ("2. Login And Session", [
        "Open the application. The Login window appears centered on the screen.",
        "Enter your username and password, then click LogIn or press Enter.",
        "The eye button beside the password field shows or hides the password.",
        "If failed-login lock is enabled in Settings, too many wrong attempts lock login until an authorized user resets the setting.",
        "After login, the dashboard opens and shows the current user, role, date/time, menu, summaries, and notifications."
    ]),
    ("3. Main Dashboard", [
        "Use the top/side menu to open Dashboard, Sales, Inventory, Supplier, Customers, Reports, Earning Report, Warehouse Reports, Users, Settings, and Logout.",
        "Only screens allowed by the Permission form are visible.",
        "Dashboard summary cards show sales/inventory/customer/supplier indicators when permitted.",
        "Notification/audit buttons show recent activity. Delete notification requires notification-delete permission.",
        "The dashboard embeds child screens inside the main content area, so most work stays in one main window."
    ]),
    ("4. Sales Workflow", [
        "Click Sales. The Start Work screen appears. Click START to create a work-history start time, then the Sales form opens.",
        "Search products by name or barcode, or choose a category and click a product button.",
        "Enter quantity before adding products. If quantity is empty, the system uses 1.",
        "Select customer, paid amount, currency (Dollar or Lebanon), and payment method.",
        "Click بيع / Save Sale. Choose whether the sale is with invoice/print or without invoice. Review the sale before final save.",
        "The review window lets you delete an item, see totals, paid amount, change due, current customer balance, and balance after sale.",
        "After saving, the system updates stock, customer balance, sale items, and audit logs."
    ]),
    ("5. Invoice Tools In Sales", [
        "Enter invoice number in the invoice field.",
        "عرض الفاتورة opens the invoice view window. The invoice includes customer, payment method, item list, totals, and balances.",
        "Use the طباعة button in the invoice window to preview/print the invoice.",
        "طبع الفاتورة excel exports the selected invoice to Excel.",
        "مسح الفاتورة deletes an invoice and its sale items when allowed.",
        "Returned registers returns. You can return a whole invoice or a single product by invoice number/barcode."
    ]),
    ("6. Inventory / Products", [
        "Inventory opens the combined product, category, and supplier management screen.",
        "Products grid supports search, edit, delete, category filtering, supplier display, stock, barcode, cost, sale prices, and dates.",
        "Use Add Product to open the product form. Required fields include product name, category, supplier, cost, sale price, stock, and optional barcode.",
        "Products with barcode can be searched/scanned in Sales. Products without barcode appear under category product buttons.",
        "Edit permission controls edit buttons. Delete permission controls delete buttons."
    ]),
    ("7. Categories", [
        "Categories group products and drive the Sales category product buttons.",
        "Add category before adding products if the category does not exist.",
        "Category delete is blocked when products still use that category.",
        "Category edit/delete are controlled by Products screen permissions."
    ]),
    ("8. Suppliers", [
        "Supplier screen manages supplier contact details and balances.",
        "Add supplier with name, contact information, and address.",
        "Supplier balance adjustment supports adding or reducing balance in USD or L.L.",
        "Supplier Balance form shows supplier list, contact information, date range, product items supplied, balances, print, and Excel export.",
        "Print in Supplier Balance opens a preview for supplier invoice/items."
    ]),
    ("9. Customers", [
        "Customers screen manages customer list, contact details, and balances.",
        "Add customer using name, phone, email, and optional opening balance.",
        "Balance + and - update the customer balance and balance date/time.",
        "Customer balances are also updated during sales depending on paid amount and sale total.",
        "Customer add/save/delete actions follow Customers screen permissions."
    ]),
    ("10. Reports", [
        "Reports screen includes report type, From Date, To Date, User filter, Refresh, daily/monthly/custom save, and Excel export.",
        "Available reports include Sales by User, Sales by Product, Daily Sales, Invoice Details, Returns, Customer Balances, Low Stock, and Work History.",
        "Work History shows User Name, ID, Work Date, Start Time, End Time, Hours, Minutes, Total Minutes, and Status.",
        "Start Time and End Time show time only; Work Date shows the date.",
        "Use Excel export to save report output for sharing or archive."
    ]),
    ("11. Earning Reports", [
        "Earning Report summarizes revenue/profit information by date range.",
        "Use filters and refresh to calculate the report.",
        "Export to Excel when you need a file copy.",
        "This screen is visible only when EarningReports permission is granted."
    ]),
    ("12. Warehouse Reports", [
        "Warehouse Reports display product stock, cost, sale price, stock value, and supplier information.",
        "Use search/filter and Refresh to update the grid.",
        "Low-stock threshold comes from Settings.",
        "Export to Excel for physical inventory checks or accounting."
    ]),
    ("13. Users", [
        "Users screen creates, edits, and deletes application users.",
        "Password rules come from Settings, including minimum expectations and maximum length.",
        "Default role for new users comes from Settings.",
        "User edit is controlled by Edit permission; create/delete by Create/Delete permissions."
    ]),
    ("14. Permissions / Options", [
        "Open Settings > Options to manage user permissions.",
        "Select a user, then grant View, Create, Edit, Save, and Delete for each screen.",
        "View controls whether a screen/menu is visible and whether the form opens.",
        "Create controls add/new buttons.",
        "Edit controls edit/update buttons and edit grid columns.",
        "Save controls save/change actions such as balance adjustments and settings save.",
        "Delete controls delete buttons and delete grid columns.",
        "Notification permission controls activity notification visibility and notification deletion."
    ]),
    ("15. Settings", [
        "Settings controls currency, language, theme, default role, password maximum length, session timeout, tax percentage, max discount, low-stock alert, report format, backup schedule, exchange rate, audit logs, refund approval, and login lock attempts.",
        "Changing language applies English/Arabic UI text and right-to-left direction for Arabic.",
        "Saving settings shows a loading overlay while the system applies settings to open forms.",
        "Reset restores default values in the Settings screen before saving."
    ]),
    ("16. Audit Log And Notifications", [
        "The system records important actions such as create, edit, delete, sales, reports, supplier/customer balance changes, and permission changes.",
        "Dashboard notification area shows recent audit entries when permission is granted.",
        "Managers can search/clear notification search and delete audit notification rows when delete notification permission is granted."
    ]),
    ("17. Printing And Exporting", [
        "Sales invoice view can print through print preview.",
        "Supplier invoice/items can print through print preview.",
        "Sales invoice, reports, earning reports, warehouse reports, and supplier invoice can export to Excel depending on the screen.",
        "Use Save As dialogs to choose file locations for exported reports."
    ]),
    ("18. Daily Recommended Use", [
        "Login with your user account.",
        "Open Sales and click START before selling.",
        "Confirm product stock and customer before saving each sale.",
        "Use invoice view/print/export when the customer needs a receipt.",
        "At the end of the day, use dashboard end-work action so Work History records the end time.",
        "Review Reports, Earning Reports, and Warehouse Reports for daily control."
    ]),
    ("19. Troubleshooting", [
        "If a menu/button is missing, check Settings > Options and confirm View/Create/Edit/Save/Delete permissions for that user.",
        "If Supplier does not open, grant View permission for Suppliers, not Products.",
        "If a product is not visible as a category button in Sales, verify it has no barcode and belongs to the selected category.",
        "If reports are empty, check date range, user filter, and whether records exist.",
        "If language changes are slow, wait until the loading overlay closes.",
        "If printing does not work, confirm the Windows printer/default printer settings."
    ])
]


arabic_sections = [
    ("١. هدف التطبيق", [
        "نظام نقاط البيع يدير المبيعات والمخزون والزبائن والموردين والأرصدة والتقارير والمستخدمين والصلاحيات والإعدادات وسجل الدوام من لوحة واحدة.",
        "التطبيق مناسب لعمل الكاشير اليومي، مراقبة المخزون، متابعة أرصدة الزبائن والموردين، واستخراج تقارير الإدارة.",
        "ظهور الشاشات والأزرار يعتمد على الصلاحيات المعطاة من Settings ثم Options."
    ]),
    ("٢. تسجيل الدخول والجلسة", [
        "افتح التطبيق. تظهر شاشة تسجيل الدخول في وسط الشاشة.",
        "اكتب اسم المستخدم وكلمة المرور ثم اضغط LogIn أو Enter.",
        "زر العين بجانب كلمة المرور لإظهار أو إخفاء كلمة المرور.",
        "إذا كان قفل النظام بعد محاولات فاشلة مفعلاً من الإعدادات، سيتم منع الدخول بعد العدد المحدد من المحاولات.",
        "بعد تسجيل الدخول تفتح لوحة التحكم وفيها اسم المستخدم والدور والوقت والقوائم والتنبيهات."
    ]),
    ("٣. لوحة التحكم الرئيسية", [
        "من القائمة يمكنك فتح Dashboard وSales وInventory وSupplier وCustomers وReports وEarning Report وWarehouse Reports وUsers وSettings وLogout.",
        "تظهر فقط الشاشات المسموحة للمستخدم من شاشة الصلاحيات.",
        "بطاقات الملخص تعرض معلومات المبيعات والمخزون وأرصدة الزبائن والموردين حسب الصلاحية.",
        "زر التنبيهات يعرض آخر العمليات المسجلة في سجل النشاط. حذف التنبيهات يحتاج صلاحية خاصة.",
        "الشاشات تفتح داخل لوحة التحكم حتى يبقى العمل في نافذة واحدة."
    ]),
    ("٤. طريقة البيع", [
        "اضغط Sales ثم START لبدء الدوام وتسجيل وقت البداية، بعدها تفتح شاشة البيع.",
        "ابحث عن المنتج بالاسم أو الباركود، أو اختر الفئة واضغط زر المنتج.",
        "اكتب الكمية قبل إضافة المنتج. إذا تركتها فارغة يستخدم النظام الكمية 1.",
        "اختر الزبون، المبلغ المدفوع، العملة Dollar أو Lebanon، وطريقة الدفع.",
        "اضغط بيع. اختر هل البيع مع فاتورة/طباعة أو بدون فاتورة، ثم راجع العملية قبل الحفظ النهائي.",
        "نافذة المراجعة تعرض الأصناف، إمكانية حذف صنف، الإجماليات، المبلغ المدفوع، المبلغ الذي يجب إرجاعه، رصيد الزبون الحالي، والرصيد بعد البيع.",
        "بعد الحفظ يقوم النظام بتحديث المخزون ورصيد الزبون وأصناف الفاتورة وسجل النشاط."
    ]),
    ("٥. أدوات الفاتورة في المبيعات", [
        "اكتب رقم الفاتورة في خانة رقم الفاتورة.",
        "زر عرض الفاتورة يفتح نافذة تحتوي بيانات الزبون وطريقة الدفع والأصناف والإجماليات والأرصدة.",
        "زر طباعة داخل نافذة الفاتورة يفتح معاينة الطباعة.",
        "زر طبع الفاتورة excel يصدر الفاتورة إلى Excel.",
        "زر مسح الفاتورة يحذف الفاتورة وأصنافها عند وجود صلاحية حذف.",
        "زر Returned يستخدم لتسجيل المرتجعات، سواء فاتورة كاملة أو منتج واحد حسب رقم الفاتورة أو الباركود."
    ]),
    ("٦. المخزون والمنتجات", [
        "شاشة Inventory تجمع إدارة المنتجات والفئات والموردين.",
        "جدول المنتجات يعرض البحث والتعديل والحذف والفئة والمورد والكمية والباركود والكلفة وسعر البيع والتاريخ.",
        "زر إضافة منتج يفتح نموذج المنتج. الحقول الأساسية: اسم المنتج، الفئة، المورد، الكلفة، سعر البيع، الكمية، والباركود إذا وجد.",
        "المنتج الذي له باركود يمكن البحث عنه أو مسحه في شاشة البيع. المنتج بدون باركود يظهر كزر داخل فئته.",
        "صلاحية Edit تتحكم بالتعديل، وصلاحية Delete تتحكم بالحذف."
    ]),
    ("٧. الفئات", [
        "الفئات تنظم المنتجات وتظهر في شاشة البيع لاختيار أزرار المنتجات.",
        "أضف الفئة قبل إضافة المنتجات إذا كانت غير موجودة.",
        "لا يمكن حذف فئة مستخدمة من منتجات موجودة.",
        "تعديل وحذف الفئات يتبع صلاحيات شاشة Products."
    ]),
    ("٨. الموردون", [
        "شاشة Supplier تدير بيانات الموردين وأرصدتهم.",
        "يمكن إضافة مورد باسم وبيانات اتصال وعنوان.",
        "تعديل رصيد المورد يدعم الزيادة أو النقصان بالدولار أو الليرة.",
        "شاشة Supplier Balance تعرض الموردين وبيانات الاتصال وفترة التاريخ والأصناف المرتبطة بالمورد والأرصدة والطباعة والتصدير إلى Excel.",
        "زر Print في Supplier Balance يفتح معاينة لطباعة فاتورة/أصناف المورد."
    ]),
    ("٩. الزبائن", [
        "شاشة Customers تدير قائمة الزبائن وبيانات الاتصال والأرصدة.",
        "إضافة زبون تتم بالاسم والهاتف والبريد ورصيد افتتاحي اختياري.",
        "أزرار + و - تعدل رصيد الزبون وتاريخ تعديل الرصيد.",
        "رصيد الزبون يتحدث أيضاً عند البيع حسب إجمالي الفاتورة والمبلغ المدفوع.",
        "الإضافة والحفظ والحذف في الزبائن تتبع صلاحيات Customers."
    ]),
    ("١٠. التقارير", [
        "شاشة Reports تحتوي نوع التقرير، من تاريخ، إلى تاريخ، فلتر المستخدم، Refresh، حفظ يومي/شهري/مخصص، وتصدير Excel.",
        "التقارير المتاحة: المبيعات حسب المستخدم، المبيعات حسب المنتج، المبيعات اليومية، تفاصيل الفواتير، المرتجعات، أرصدة الزبائن، المخزون المنخفض، وسجل الدوام.",
        "Work History يعرض اسم المستخدم، الرقم، تاريخ العمل، وقت البداية، وقت النهاية، الساعات، الدقائق، مجموع الدقائق، والحالة.",
        "Start Time وEnd Time يعرضان الوقت فقط، وWork Date يعرض التاريخ فقط.",
        "استخدم Excel Export لحفظ نسخة من التقرير."
    ]),
    ("١١. تقارير الأرباح", [
        "Earning Report يعرض ملخص الإيرادات والأرباح حسب الفترة.",
        "استخدم الفلاتر والتحديث لحساب التقرير.",
        "يمكن التصدير إلى Excel عند الحاجة.",
        "تظهر الشاشة فقط لمن لديه صلاحية EarningReports."
    ]),
    ("١٢. تقارير المستودع", [
        "Warehouse Reports تعرض المنتجات والكمية والكلفة وسعر البيع وقيمة المخزون والمورد.",
        "استخدم البحث أو الفلتر ثم Refresh لتحديث الجدول.",
        "حد المخزون المنخفض يأتي من Settings.",
        "التصدير إلى Excel مفيد للجرد والمحاسبة."
    ]),
    ("١٣. المستخدمون", [
        "شاشة Users تستخدم لإنشاء وتعديل وحذف المستخدمين.",
        "قواعد كلمة المرور تأتي من Settings مثل الحد الأقصى للطول.",
        "الدور الافتراضي للمستخدم الجديد يأتي من Settings.",
        "تعديل المستخدم يحتاج صلاحية Edit، والإضافة والحذف يحتاجان Create وDelete."
    ]),
    ("١٤. الصلاحيات / Options", [
        "افتح Settings ثم Options لإدارة صلاحيات المستخدمين.",
        "اختر المستخدم ثم حدد View وCreate وEdit وSave وDelete لكل شاشة.",
        "View يحدد ظهور الشاشة وإمكانية فتحها.",
        "Create يتحكم بأزرار الإضافة.",
        "Edit يتحكم بأزرار وأعمدة التعديل.",
        "Save يتحكم بعمليات الحفظ مثل تعديل الأرصدة وحفظ الإعدادات.",
        "Delete يتحكم بأزرار وأعمدة الحذف.",
        "صلاحيات التنبيهات تتحكم بظهور سجل النشاط وحذف التنبيهات."
    ]),
    ("١٥. الإعدادات", [
        "Settings تتحكم بالعملة واللغة والثيم والدور الافتراضي وطول كلمة المرور ومدة الجلسة والضريبة والخصم وحد المخزون المنخفض وصيغة التقارير والنسخ الاحتياطي وسعر الصرف وسجل النشاط وموافقة المرتجع وقفل الدخول.",
        "تغيير اللغة يطبق English أو Arabic ويغير اتجاه الواجهة للعربية.",
        "عند حفظ الإعدادات تظهر شاشة Loading أثناء تطبيق التغييرات على النوافذ المفتوحة.",
        "Reset يعيد القيم الافتراضية على الشاشة قبل الحفظ."
    ]),
    ("١٦. سجل النشاط والتنبيهات", [
        "النظام يسجل العمليات المهمة مثل الإضافة والتعديل والحذف والمبيعات والتقارير وتعديل الأرصدة والصلاحيات.",
        "منطقة التنبيهات في Dashboard تعرض آخر السجلات عند وجود الصلاحية.",
        "يمكن البحث في التنبيهات وحذفها عند وجود صلاحية حذف التنبيهات."
    ]),
    ("١٧. الطباعة والتصدير", [
        "فاتورة المبيعات يمكن طباعتها من نافذة عرض الفاتورة.",
        "فاتورة/أصناف المورد يمكن طباعتها من Supplier Balance.",
        "يمكن تصدير الفواتير والتقارير وتقارير الأرباح وتقارير المستودع وفاتورة المورد إلى Excel حسب الشاشة.",
        "اختر مكان الحفظ من نافذة Save As عند التصدير."
    ]),
    ("١٨. الاستخدام اليومي المقترح", [
        "سجل الدخول بحسابك.",
        "افتح Sales واضغط START قبل البيع.",
        "تأكد من المنتج والكمية والزبون قبل حفظ كل عملية.",
        "استخدم عرض الفاتورة أو الطباعة أو التصدير عندما يحتاج الزبون إيصالاً.",
        "في نهاية اليوم استخدم إنهاء الدوام من Dashboard حتى يسجل النظام وقت النهاية.",
        "راجع Reports وEarning Reports وWarehouse Reports لمراقبة العمل اليومي."
    ]),
    ("١٩. حل المشاكل", [
        "إذا كان زر أو شاشة غير ظاهرة، افتح Settings > Options وتأكد من صلاحيات View/Create/Edit/Save/Delete للمستخدم.",
        "إذا لم تفتح شاشة Supplier، أعط المستخدم View على Suppliers وليس Products.",
        "إذا لم يظهر المنتج كزر في شاشة البيع، تأكد أنه بدون باركود ومربوط بالفئة المختارة.",
        "إذا كان التقرير فارغاً، تأكد من التاريخ وفلتر المستخدم ووجود بيانات.",
        "إذا استغرق تغيير اللغة وقتاً، انتظر حتى تختفي شاشة Loading.",
        "إذا لم تعمل الطباعة، تأكد من إعدادات الطابعة الافتراضية في Windows."
    ])
]


screen_rows_en = [
    ("Dashboard", "Overview, metrics, notifications, activity, menu access"),
    ("Sales", "Start work, sell products, invoices, print, return, delete invoice"),
    ("Products / Inventory", "Products, categories, suppliers, stock and barcode control"),
    ("Suppliers", "Supplier details, balances, item history, print/export"),
    ("Customers", "Customer details and balance updates"),
    ("Reports", "Sales, products, invoices, returns, balances, low stock, work history"),
    ("Earning Reports", "Revenue/profit summaries and Excel export"),
    ("Warehouse", "Stock and warehouse value reports"),
    ("Users", "Create/edit/delete application users"),
    ("Settings", "Currency, language, theme, security, exchange rate, defaults"),
    ("Options", "Permission management")
]

screen_rows_ar = [
    ("Dashboard", "الملخصات والتنبيهات وسجل النشاط والقوائم"),
    ("Sales", "بدء الدوام والبيع والفواتير والطباعة والمرتجعات والحذف"),
    ("Products / Inventory", "المنتجات والفئات والموردون والمخزون والباركود"),
    ("Suppliers", "بيانات الموردين والأرصدة والتاريخ والطباعة والتصدير"),
    ("Customers", "بيانات الزبائن وتعديل الأرصدة"),
    ("Reports", "المبيعات والمنتجات والفواتير والمرتجعات والأرصدة والمخزون وسجل الدوام"),
    ("Earning Reports", "ملخص الإيرادات والأرباح والتصدير"),
    ("Warehouse", "تقارير المخزون وقيمة المستودع"),
    ("Users", "إنشاء وتعديل وحذف المستخدمين"),
    ("Settings", "العملة واللغة والثيم والأمان وسعر الصرف والإعدادات الافتراضية"),
    ("Options", "إدارة الصلاحيات")
]


def build_manual(path, title, subtitle, sections, screen_rows, rtl=False):
    doc = setup_document(rtl)
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.CENTER
    if rtl:
        set_paragraph_rtl(title_p)
    run = title_p.add_run(title)
    set_run_font(run, "Tahoma" if rtl else "Segoe UI", 24, True, (15, 23, 42))
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.CENTER
    if rtl:
        set_paragraph_rtl(sub_p)
    run = sub_p.add_run(subtitle)
    set_run_font(run, "Tahoma" if rtl else "Segoe UI", 12, False, (71, 85, 105))
    doc.add_paragraph()

    add_heading(doc, "Quick Screen Guide" if not rtl else "دليل سريع للشاشات", 1, rtl)
    add_table(
        doc,
        ["Screen", "Purpose"] if not rtl else ["الشاشة", "الاستخدام"],
        screen_rows,
        rtl,
    )

    add_heading(doc, "Contents" if not rtl else "المحتويات", 1, rtl)
    add_bullets(doc, [heading for heading, _ in sections], rtl)
    doc.add_page_break()

    for heading, bullets in sections:
        add_heading(doc, heading, 1, rtl)
        add_bullets(doc, bullets, rtl)
        doc.add_paragraph()

    add_heading(doc, "Permission Matrix Summary" if not rtl else "ملخص الصلاحيات", 1, rtl)
    rows = [
        ("View", "Open and see the screen") if not rtl else ("View", "فتح الشاشة ورؤيتها"),
        ("Create", "Use add/new actions") if not rtl else ("Create", "استخدام أزرار الإضافة"),
        ("Edit", "Use edit/update actions") if not rtl else ("Edit", "استخدام أزرار التعديل"),
        ("Save", "Use save/balance/settings actions") if not rtl else ("Save", "استخدام الحفظ وتعديل الأرصدة والإعدادات"),
        ("Delete", "Use delete actions") if not rtl else ("Delete", "استخدام الحذف"),
    ]
    add_table(doc, ["Permission", "Meaning"] if not rtl else ["الصلاحية", "المعنى"], rows, rtl)

    add_heading(doc, "Final Checklist" if not rtl else "قائمة تحقق أخيرة", 1, rtl)
    checklist = [
        "Confirm users and roles are correct.",
        "Grant permissions from Settings > Options.",
        "Set exchange rate before sales.",
        "Set low-stock threshold before warehouse reporting.",
        "Start work before daily sales and end work at day close.",
        "Export or print reports/invoices when needed."
    ] if not rtl else [
        "تأكد من صحة المستخدمين والأدوار.",
        "أعط الصلاحيات من Settings > Options.",
        "حدد سعر الصرف قبل البيع.",
        "حدد حد المخزون المنخفض قبل تقارير المستودع.",
        "ابدأ الدوام قبل البيع اليومي وأنهِ الدوام في آخر اليوم.",
        "اطبع أو صدّر التقارير والفواتير عند الحاجة."
    ]
    add_numbered(doc, checklist, rtl)

    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("POS System User Manual" if not rtl else "دليل استخدام نظام نقاط البيع")
    set_run_font(run, "Tahoma" if rtl else "Segoe UI", 9, False, (100, 116, 139))
    doc.save(path)


if __name__ == "__main__":
    build_manual(
        OUTPUT_DIR / "POS_System_User_Manual_English.docx",
        "POS System User Manual",
        "Complete guide for using the POS application",
        english_sections,
        screen_rows_en,
        rtl=False,
    )

    build_manual(
        OUTPUT_DIR / "POS_System_User_Manual_Arabic.docx",
        "دليل استخدام نظام نقاط البيع",
        "شرح كامل لاستخدام التطبيق",
        arabic_sections,
        screen_rows_ar,
        rtl=True,
    )

    print(OUTPUT_DIR)
